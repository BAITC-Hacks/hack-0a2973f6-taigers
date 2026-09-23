import base64
import io
from pathlib import Path

import fitz
from docx import Document
from fastapi import UploadFile
from openpyxl import load_workbook

from app.config import Settings


class FileExtractionError(ValueError):
    pass


async def extract_upload(file: UploadFile, settings: Settings) -> dict:
    content = await file.read()
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise FileExtractionError(f"File exceeds {settings.max_upload_mb} MB")
    suffix = Path(file.filename or "").suffix.casefold()
    mime = file.content_type or "application/octet-stream"
    if suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        encoded = base64.b64encode(content).decode("ascii")
        return {"kind": "image", "image_data_url": f"data:{mime};base64,{encoded}", "text": ""}
    if suffix == ".pdf":
        document = fitz.open(stream=content, filetype="pdf")
        return {"kind": "document", "text": _limit("\n".join(page.get_text("text") for page in document))}
    if suffix == ".docx":
        document = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        return {"kind": "document", "text": _limit(text)}
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets[:5]:
            rows.append(f"Лист: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if values:
                    rows.append(" | ".join(values))
                if len(rows) >= 500:
                    break
        return {"kind": "document", "text": _limit("\n".join(rows))}
    raise FileExtractionError("Supported formats: PDF, DOCX, XLSX, JPEG, PNG, WEBP")


def _limit(text: str, length: int = 20_000) -> str:
    clean = text.strip()
    return clean[:length] + ("\n[текст сокращён]" if len(clean) > length else "")
